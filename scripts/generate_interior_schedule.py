import csv
import os

output_path = r"c:\Users\SEOP\Desktop\seop architecture\docs\interior_finish_schedule.csv"

# Data structure: Type, Room, Element, Material, Spec, Quantity(m2/unit), Unit, Unit Price(KRW), Total Price(KRW), Remarks
data = [
    # Header
    ["타입(Type)", "실명(Room)", "부위(Element)", "마감재(Material)", "규격/상세(Spec)", "수량", "단위", "단가(원)", "합계(원)", "비고"],
    
    # 59 Type (Modern Chic)
    ["59㎡", "거실/주방", "바닥", "강마루 (Oak)", "115 x 800 x 7.5t", 45, "m2", 45000, 2025000, "광폭 강마루 (자재비)"],
    ["59㎡", "거실/주방", "벽", "실크벽지", "LG 지인/개나리", 110, "m2", 5000, 550000, "라이트 그레이 톤"],
    ["59㎡", "거실/주방", "천정", "실크도배", "우물천정 포함", 45, "m2", 5000, 225000, "간접조명 별도"],
    ["59㎡", "침실1(안방)", "바닥", "강마루 (Oak)", "동일 스펙", 18, "m2", 45000, 810000, "-"],
    ["59㎡", "침실1(안방)", "벽", "실크벽지", "웜 그레이", 45, "m2", 5000, 225000, "-"],
    ["59㎡", "공용욕실", "바닥/벽", "포세린 타일", "600 x 600", 28, "m2", 35000, 980000, "논슬립 (자재비)"],
    
    # 84 Type (Urban Luxury)
    ["84㎡", "거실/주방", "바닥", "수입 원목마루", "190 x 1900 x 12t", 65, "m2", 120000, 7800000, "브러쉬 텍스처"],
    ["84㎡", "거실/주방", "벽", "친환경 도장(벤자민무어)", "스커프엑스", 150, "m2", 25000, 3750000, "자재비 기준"],
    ["84㎡", "주방 상판", "가구", "엔지니어드 스톤", "칸스톤/비아테라", 5.5, "m", 350000, 1925000, "아일랜드 포함 (m당 단가)"],
    ["84㎡", "침실1(안방)", "바닥", "원목마루", "동일 스펙", 24, "m2", 120000, 2880000, "-"],
    ["84㎡", "부부욕실", "바닥/벽", "이태리 타일", "600 x 1200", 35, "m2", 70000, 2450000, "호텔식 조적 욕조"],
    
    # 114 Type (High-end Premium)
    ["114㎡", "거실/주방", "바닥", "천연 대리석", "1200 x 1200", 85, "m2", 150000, 12750000, "보티치노/크레마마필"],
    ["114㎡", "거실", "아트월", "박판 세라믹", "3200 x 1600", 15, "m2", 180000, 2700000, "북매치 시공"],
    ["114㎡", "거실/복도", "벽", "무늬목 패널", "천연 건식 무늬목", 180, "m2", 150000, 27000000, "히든도어 포함"],
    ["114㎡", "주방", "가구", "수입 주방 가구", "Bulthaup/Arclinea 급", 1, "식", 85000000, 85000000, "빌트인 가전 별도"],
    ["114㎡", "마스터존", "바닥", "광폭 원목마루", "240 x 2200", 40, "m2", 150000, 6000000, "헤링본 패턴"],
    ["114㎡", "마스터 욕실", "전체", "천연석/대형타일", "-", 45, "m2", 120000, 5400000, "독립형 욕조, 수전 특화"]
]

try:
    # Ensure directory exists
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    # Change extension to .xlsx
    excel_output_path = output_path.replace('.csv', '.xlsx')
    
    # Use openpyxl for real Excel file generation
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    
    wb = Workbook()
    ws = wb.active
    ws.title = "실내마감상세"
    
    # Add data
    for row in data:
        ws.append(row)
        
    # Styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Apply style to Header (Row 1)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Apply style to Data
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = thin_border
            # Center align for first few columns, right align for numbers
            if cell.column <= 5 or cell.column == 7: # Type, Room, Element, Material, Spec, Unit
                cell.alignment = center_align
            elif cell.column == 6 or cell.column >= 8: # Quantity, Price, Total
                 cell.number_format = '#,##0' # Number format

    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2) * 1.1
        ws.column_dimensions[column].width = adjusted_width

    wb.save(excel_output_path)
    print(f"Successfully created interior schedule at: {excel_output_path}")

except Exception as e:
    print(f"Error: {e}")
